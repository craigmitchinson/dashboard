"""Converts docs/FEATURE-CATALOGUE.md into a styled Word document.

    node tools/build-feature-catalogue.mjs      # regenerate the markdown first
    python tools/build-feature-catalogue-docx.py

Requires python-docx (pip install python-docx). Output:
docs/Feature Catalogue - IA Performance Dashboard.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "FEATURE-CATALOGUE.md"
OUT = ROOT / "docs" / "Feature Catalogue - IA Performance Dashboard.docx"
TEAL = RGBColor(0x0B, 0x32, 0x39)
RED = RGBColor(0xC8, 0x1E, 0x2B)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(par, text):
    for t in INLINE.split(text):
        if not t:
            continue
        if t.startswith("**") and t.endswith("**"):
            par.add_run(t[2:-2]).bold = True
        elif t.startswith("`") and t.endswith("`"):
            r = par.add_run(t[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = TEAL
        elif t.startswith("*") and t.endswith("*"):
            par.add_run(t[1:-1]).italic = True
        else:
            par.add_run(t)


def shade(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def split_row(line):
    return [c.strip().replace("\\|", "|") for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    headers = split_row(rows[0])
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h.upper())
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = TEAL
        shade(cell, "F3EFE8")
    for raw in rows[2:]:
        values = split_row(raw)
        cells = table.add_row().cells
        for i in range(len(headers)):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_inline(p, values[i] if i < len(values) else "")
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph()


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 13.5)):
        st = styles[name]
        st.font.name = "Georgia" if name in ("Title", "Heading 1") else "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = TEAL
        st.element.rPr.rFonts.set(qn("w:eastAsia"), st.font.name)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    subtitle_done = False
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("<!--") or ln.strip() in ("", "---"):
            i += 1
            continue
        if ln.startswith("# "):
            doc.add_paragraph(ln[2:], style="Title")
        elif ln.startswith("## "):
            title = ln[3:]
            if not subtitle_done:
                p = doc.add_paragraph()
                r = p.add_run(title)
                r.font.size = Pt(13)
                r.font.color.rgb = RED
                r.bold = True
                subtitle_done = True
            else:
                if title.startswith("1. "):
                    doc.add_page_break()
                doc.add_heading(title, level=1)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:], level=2)
        elif ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        elif ln.startswith("- "):
            while i < len(lines) and lines[i].startswith("- "):
                add_inline(doc.add_paragraph(style="List Bullet"), lines[i][2:])
                i += 1
            continue
        else:
            p = doc.add_paragraph()
            add_inline(p, ln)
            p.paragraph_format.space_after = Pt(6)
        i += 1

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Intelligent Automation Performance Dashboard · Feature catalogue v1.0 · Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = TEAL
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    rr = OxmlElement("w:r")
    tt = OxmlElement("w:t")
    tt.text = "1"
    rr.append(tt)
    fld.append(rr)
    footer._p.append(fld)

    doc.save(OUT)
    print(f"written: {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
